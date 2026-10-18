import json

from django.db.models import Q, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views import View
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from apps.academic.models import AcademicYearModel
from apps.academic.services.active_year import get_active_year
from apps.finance.forms import (
    SalaryPaymentCreateForm,
    StaffSalaryCreateForm,
    StaffSalaryUpdateForm,
)
from apps.finance.models import StaffSalaryModel
from apps.teacher.models import TeacherModel


class StaffSalaryCreateView(View):
    template_name = "finance/salary/create.html"

    def get(self, request):
        form = StaffSalaryCreateForm()
        return render(request, self.template_name, {"form": form})

    def post(self, request):
        form = StaffSalaryCreateForm(request.POST)

        if form.is_valid():
            salary = form.save()
            print("SALARY SAVED:", salary.pk)
            return redirect("finance:salary_list")

        print("FORM ERRORS:", form.errors)
        print("FORM NON FIELD ERRORS:", form.non_field_errors())
        print("POST DATA:", request.POST.dict())

        return render(request, self.template_name, {"form": form})


class StaffSalaryListView(View):
    template_name = "finance/salary/list.html"

    def get(self, request):
        year_id = request.GET.get("year")
        search = request.GET.get("search")

        salaries = StaffSalaryModel.objects.select_related("teacher", "academic_year")

        # فیلتر سال تحصیلی
        if year_id:
            salaries = salaries.filter(academic_year_id=year_id)

        # جستجو معلم
        if search:
            salaries = salaries.filter(
                Q(teacher__user__first_name__icontains=search)
                | Q(teacher__user__last_name__icontains=search)
                | Q(teacher__national_id__icontains=search)
            )

        years = AcademicYearModel.objects.all()

        return render(
            request,
            self.template_name,
            {
                "salaries": salaries,
                "years": years,
            },
        )


class StaffSalaryUpdateView(View):
    template_name = "finance/salary/update.html"

    def get(self, request, salary_id):
        salary = get_object_or_404(StaffSalaryModel, id=salary_id)
        form = StaffSalaryUpdateForm(instance=salary)
        return render(request, self.template_name, {"form": form, "salary": salary})

    def post(self, request, salary_id):
        salary = get_object_or_404(StaffSalaryModel, id=salary_id)
        form = StaffSalaryUpdateForm(request.POST, instance=salary)

        if form.is_valid():
            form.save()
            return redirect("finance:salary_list")

        return render(request, self.template_name, {"form": form, "salary": salary})


class MonthlySalaryPrepareView(View):
    template_name = "finance/salary/monthly_prepare.html"

    def get(self, request):
        active_year = get_active_year(request)

        salaries = (
            StaffSalaryModel.objects.select_related("teacher", "academic_year").filter(
                academic_year=active_year
            )
            if active_year
            else []
        )

        return render(
            request,
            self.template_name,
            {
                "active_year": active_year,
                "salaries": salaries,
            },
        )


class MonthlySalaryExportView(View):

    def post(self, request):
        active_year = get_active_year(request)
        month = request.POST.get("month")
        payload = json.loads(request.POST.get("payload") or "[]")

        wb = Workbook()
        ws = wb.active
        ws.title = "لیست پرداخت حقوق"

        # استایل‌ها
        header_font = Font(bold=True, size=12)
        center = Alignment(horizontal="center", vertical="center")
        fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # عنوان بزرگ
        ws.merge_cells("A1:E1")
        ws["A1"] = f"لیست پرداخت حقوق معلمان - ماه {month} - سال {active_year.name}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = center

        # هدر جدول
        headers = ["نام معلم", "کد ملی", "سال تحصیلی", "ماه", "مبلغ پرداختی"]
        ws.append(headers)

        for col in "ABCDE":
            cell = ws[f"{col}2"]
            cell.font = header_font
            cell.alignment = center
            cell.fill = fill
            cell.border = thin_border

        # داده‌ها
        row = 3
        for item in payload:
            teacher = get_object_or_404(TeacherModel, id=item["teacher_id"])
            ws.append(
                [
                    teacher.user.full_name,
                    teacher.national_id,
                    active_year.name,
                    month,
                    item["amount"],
                ]
            )

            # استایل‌دهی ردیف
            for col in "ABCDE":
                cell = ws[f"{col}{row}"]
                cell.alignment = center
                cell.border = thin_border

            row += 1

        # تنظیم عرض ستون‌ها
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 10
        ws.column_dimensions["E"].width = 18

        # فریز کردن ردیف هدر
        ws.freeze_panes = "A3"

        # خروجی اکسل
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = f"salary_{active_year.name}_month_{month}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        wb.save(response)

        return response


class SalaryPaymentCreateView(View):
    template_name = "finance/salary/payment_create.html"

    def get(self, request, salary_id):
        salary = get_object_or_404(StaffSalaryModel, id=salary_id)
        form = SalaryPaymentCreateForm(salary=salary)
        return render(request, self.template_name, {"form": form, "salary": salary})

    def post(self, request, salary_id):
        salary = get_object_or_404(StaffSalaryModel, id=salary_id)
        form = SalaryPaymentCreateForm(request.POST, salary=salary)

        if form.is_valid():
            payment = form.save(commit=False)
            payment.salary = salary
            payment.save()
            return redirect("finance:salary_list")

        return render(request, self.template_name, {"form": form, "salary": salary})


class StaffSalaryDetailView(View):
    template_name = "finance/salary/detail.html"

    def get(self, request, salary_id):
        salary = get_object_or_404(StaffSalaryModel, id=salary_id)

        # سال فعال
        active_year = get_active_year(request)

        # سال انتخاب‌شده توسط کاربر
        year_id = request.GET.get("year")

        years = AcademicYearModel.objects.all()

        # پرداخت‌ها
        payments = salary.payments.all()

        # اگر کاربر سال انتخاب کرده باشد → همان را اعمال کن
        if year_id:
            payments = payments.filter(salary__academic_year_id=year_id)
            selected_year = AcademicYearModel.objects.get(id=year_id)

        # اگر کاربر سال انتخاب نکرده باشد → سال فعال را اعمال کن
        else:
            payments = payments.filter(salary__academic_year=active_year)
            selected_year = active_year

        # مجموع پرداخت‌ها
        total_paid = payments.aggregate(total=Sum("amount"))["total"] or 0

        return render(
            request,
            self.template_name,
            {
                "salary": salary,
                "payments": payments,
                "years": years,
                "selected_year": selected_year,
                "total_paid": total_paid,
            },
        )


class TeacherSalaryDashboardView(View):
    template_name = "teacher/salary.html"

    def get(self, request):
        # معلم لاگین‌شده
        teacher = get_object_or_404(TeacherModel, user=request.user)

        # سال فعال
        active_year = get_active_year(request)

        # سال انتخاب‌شده توسط کاربر
        year_id = request.GET.get("year")

        years = AcademicYearModel.objects.all()

        # حکم‌های معلم
        salaries = StaffSalaryModel.objects.filter(teacher=teacher)

        # انتخاب سال
        if year_id:
            selected_year = AcademicYearModel.objects.get(id=year_id)
        else:
            selected_year = active_year

        # حکم سال انتخاب‌شده
        salary = salaries.filter(academic_year=selected_year).first()

        payments = []
        total_paid = 0
        remaining = 0

        if salary:
            payments = salary.payments.all()
            total_paid = payments.aggregate(total=Sum("amount"))["total"] or 0
            remaining = salary.base_amount - total_paid

        return render(
            request,
            self.template_name,
            {
                "teacher": teacher,
                "salary": salary,
                "payments": payments,
                "years": years,
                "selected_year": selected_year,
                "total_paid": total_paid,
            },
        )
